"""Regression coverage for untrusted reservation content in downloaded reports."""
from datetime import datetime, time, timezone as dt_timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase, override_settings
from openpyxl import load_workbook
from reportlab.platypus import Paragraph

from calendario.exporters import ExcelExportStrategy, PDFExportStrategy


@override_settings(TIME_ZONE='America/Argentina/Buenos_Aires', USE_TZ=True)
class ExportSecurityTests(SimpleTestCase):
    def setUp(self):
        self.reserva = SimpleNamespace(
            fecha_inicio=datetime(2026, 10, 5, 1, 30, tzinfo=dt_timezone.utc),
            fecha_fin=datetime(2026, 10, 5, 2, 30, tzinfo=dt_timezone.utc),
            titulo='Revisión <equipo> & coordinación',
            descripcion='Texto normal',
            recurso=SimpleNamespace(nombre='Sala & reuniones', capacidad=10),
            usuario=SimpleNamespace(username='operador', email='operador@example.test'),
        )
        self.start = datetime(2026, 10, 1)
        self.end = datetime(2026, 10, 31)
        self.options = {'include_descriptions': True}

    def test_pdf_treats_user_markup_as_text_without_loading_images(self):
        payload = '<img src="https://example.invalid/private.png"/>'
        self.reserva.descripcion = payload
        paragraphs = []

        def recording_paragraph(text, *args, **kwargs):
            paragraphs.append(text)
            return Paragraph(text, *args, **kwargs)

        with patch('reportlab.lib.utils.ImageReader') as image_reader, patch(
            'reportlab.platypus.Paragraph', side_effect=recording_paragraph
        ):
            response = PDFExportStrategy().export(
                [self.reserva], self.start, self.end, self.options
            )

        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertTrue(response.content.startswith(b'%PDF-'))
        image_reader.assert_not_called()
        self.assertIn('&lt;img src="https://example.invalid/private.png"/&gt;', paragraphs)
        self.assertIn('Revisión &lt;equipo&gt; &amp; coordinación', paragraphs)
        self.assertIn('Sala &amp; reuniones', paragraphs)
        self.assertIn('<b>04/10/2026</b>', paragraphs)
        self.assertIn('22:30 - 23:30', paragraphs)

    def test_excel_keeps_formulas_literal_in_details_summary_and_analysis(self):
        payload = '=HYPERLINK("https://example.invalid/", "abrir")'
        self.reserva.titulo = payload
        self.reserva.descripcion = '=1+1'
        self.reserva.recurso.nombre = '=2+2'
        self.reserva.usuario.username = '=3+3'
        self.reserva.usuario.email = '=4+4'

        response = ExcelExportStrategy().export(
            [self.reserva], self.start, self.end, self.options
        )
        self.assertEqual(response.status_code, 200, response.content[:200])
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        details = workbook['Detalle de Reservas']
        self.assertEqual(details['H2'].value, payload)
        self.assertEqual(details['L2'].value, '=1+1')
        self.assertEqual(workbook['Resumen Ejecutivo']['B11'].value, '=2+2')
        self.assertEqual(workbook['Análisis por Sala']['A4'].value, '=2+2')
        for sheet in workbook:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, 'f', (sheet.title, cell.coordinate))
        self.assertEqual(details['A2'].value.date(), datetime(2026, 10, 4).date())
        self.assertEqual(details['B2'].value, 'Domingo')
        self.assertEqual(details['C2'].value, time(22, 30))
        self.assertEqual(details['D2'].value, time(23, 30))
        self.assertEqual(details['E2'].value, 1)

    def test_export_failure_does_not_disclose_internal_exception(self):
        with patch('openpyxl.Workbook', side_effect=RuntimeError('internal/private/path')):
            response = ExcelExportStrategy().export(
                [self.reserva], self.start, self.end, self.options
            )
        self.assertEqual(response.status_code, 500)
        self.assertNotIn(b'internal/private/path', response.content)

    def test_pdf_handles_description_that_spans_multiple_pages(self):
        self.reserva.descripcion = 'Detalle de reunión con el equipo. ' * 300
        response = PDFExportStrategy().export(
            [self.reserva], self.start, self.end, self.options
        )
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertTrue(response.content.startswith(b'%PDF-'))
