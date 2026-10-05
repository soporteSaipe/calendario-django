from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.management import call_command, CommandError
from django.test import TestCase
from openpyxl import Workbook


class UserImportTests(TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name) / 'test_accounts.xlsx'
        self.output = StringIO()
        self.errors = StringIO()

    def write_rows(self, rows):
        workbook = Workbook()
        workbook.active.append(['usuario', 'contraseña'])
        for row in rows:
            workbook.active.append(row)
        workbook.save(self.path)
        workbook.close()

    def run_import(self, **options):
        call_command('cargar_usuarios_excel', archivo=str(self.path),
                     stdout=self.output, stderr=self.errors, **options)

    def test_default_import_creates_new_users_and_preserves_existing_admin(self):
        admin = User.objects.create_superuser('existing', email='original@example.test', password='Original!Key-478')
        original_hash = admin.password
        self.write_rows([['existing', 'Replacement!Key-935'], ['new-person', 'Good!Pass-684-B']])
        self.run_import()
        admin.refresh_from_db()
        self.assertEqual(admin.password, original_hash)
        self.assertEqual(admin.email, 'original@example.test')
        self.assertTrue(admin.is_staff and admin.is_superuser)
        new_user = User.objects.get(username='new-person')
        self.assertTrue(new_user.check_password('Good!Pass-684-B'))
        self.assertFalse(new_user.is_staff or new_user.is_superuser)
        self.assertNotIn('Replacement!Key-935', self.output.getvalue()+self.errors.getvalue())

    def test_existing_password_update_requires_explicit_flag_and_preserves_profile(self):
        user = User.objects.create_superuser('existing', email='original@example.test', password='Original!Key-478')
        user.first_name = 'Persona'
        user.save()
        self.write_rows([['existing', 'Replacement!Key-935']])
        self.run_import(actualizar_existentes=True)
        user.refresh_from_db()
        self.assertTrue(user.check_password('Replacement!Key-935'))
        self.assertEqual(user.email, 'original@example.test')
        self.assertEqual(user.first_name, 'Persona')
        self.assertTrue(user.is_staff and user.is_superuser)

    def test_weak_password_and_invalid_username_do_not_create_partial_accounts(self):
        self.write_rows([['weak-user', '123'], ['invalid name', 'Good!Pass-684-B'],
                         ['valid-person', 'Good!Pass-739-Z']])
        with self.assertRaises(CommandError):
            self.run_import()
        self.assertFalse(User.objects.filter(username='weak-user').exists())
        self.assertFalse(User.objects.filter(username='invalid name').exists())
        self.assertTrue(User.objects.get(username='valid-person').check_password('Good!Pass-739-Z'))

    def test_weak_update_preserves_existing_password(self):
        user = User.objects.create_user('existing', password='Original!Key-478')
        original_hash = user.password
        self.write_rows([['existing', '123']])
        with self.assertRaises(CommandError):
            self.run_import(actualizar_existentes=True)
        user.refresh_from_db()
        self.assertEqual(user.password, original_hash)

    def test_workbook_is_closed_if_iteration_fails(self):
        self.write_rows([])
        with patch('calendario.management.commands.cargar_usuarios_excel.load_workbook') as loader:
            loader.return_value.active.iter_rows.side_effect = RuntimeError('broken workbook')
            with self.assertRaises(RuntimeError):
                self.run_import()
            loader.return_value.close.assert_called_once()
