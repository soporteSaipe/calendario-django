"""
Comando de Django para cargar usuarios desde un archivo Excel
"""
import os
from django.core.management.base import BaseCommand
from django.contrib.auth.models import User
from openpyxl import load_workbook


class Command(BaseCommand):
    help = 'Carga usuarios desde un archivo Excel con columnas: usuario y contraseña'

    def add_arguments(self, parser):
        parser.add_argument(
            '--archivo',
            type=str,
            default='credenciales.xlsx',
            help='Ruta al archivo Excel (por defecto: credenciales.xlsx en raíz del proyecto)'
        )

    def handle(self, *args, **options):
        archivo_excel = options['archivo']
        
        # Verificar si el archivo existe
        if not os.path.exists(archivo_excel):
            self.stdout.write(
                self.style.WARNING(f'El archivo {archivo_excel} no existe. Saltando creación de usuarios.')
            )
            return
        
        try:
            # Cargar el archivo Excel
            self.stdout.write(f'Cargando archivo: {archivo_excel}')
            workbook = load_workbook(archivo_excel, read_only=True)
            sheet = workbook.active
            
            # Contadores
            usuarios_creados = 0
            usuarios_existentes = 0
            errores = 0
            
            # Leer las filas (asumiendo que la primera fila son encabezados)
            for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    # Verificar que la fila tenga al menos 2 columnas
                    if not row or len(row) < 2:
                        self.stdout.write(
                            self.style.WARNING(f'Fila {row_num}: Fila vacía o incompleta, saltando...')
                        )
                        continue
                    
                    usuario = row[0]
                    contraseña = row[1]
                    
                    # Validar que no estén vacíos
                    if not usuario or not contraseña:
                        self.stdout.write(
                            self.style.WARNING(f'Fila {row_num}: Usuario o contraseña vacío, saltando...')
                        )
                        errores += 1
                        continue
                    
                    # Convertir a string y limpiar espacios
                    usuario = str(usuario).strip()
                    contraseña = str(contraseña).strip()
                    
                    # Crear email
                    email = f'{usuario}@saipe.com.ar'
                    
                    # Verificar si el usuario ya existe
                    if User.objects.filter(username=usuario).exists():
                        self.stdout.write(
                            self.style.WARNING(f'Fila {row_num}: Usuario "{usuario}" ya existe, saltando...')
                        )
                        usuarios_existentes += 1
                        continue
                    
                    # Crear el usuario
                    User.objects.create_user(
                        username=usuario,
                        email=email,
                        password=contraseña,
                        is_staff=False,
                        is_superuser=False
                    )
                    
                    self.stdout.write(
                        self.style.SUCCESS(f'Fila {row_num}: Usuario "{usuario}" creado exitosamente')
                    )
                    usuarios_creados += 1
                    
                except Exception as e:
                    self.stdout.write(
                        self.style.ERROR(f'Fila {row_num}: Error al procesar - {str(e)}')
                    )
                    errores += 1
            
            # Cerrar el workbook
            workbook.close()
            
            # Resumen
            self.stdout.write('\n' + '='*60)
            self.stdout.write(self.style.SUCCESS(f'✓ Usuarios creados: {usuarios_creados}'))
            if usuarios_existentes > 0:
                self.stdout.write(self.style.WARNING(f'⚠ Usuarios que ya existían: {usuarios_existentes}'))
            if errores > 0:
                self.stdout.write(self.style.ERROR(f'✗ Errores: {errores}'))
            self.stdout.write('='*60 + '\n')
            
        except Exception as e:
            self.stdout.write(
                self.style.ERROR(f'Error al cargar el archivo: {str(e)}')
            )
            raise

